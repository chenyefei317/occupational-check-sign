import datetime
import io
import os
import re
import zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor, Pt
import pandas as pd
import qrcode
from PIL import Image
import requests
import streamlit as st
from streamlit_drawable_canvas import st_canvas

# ================= 1. 页面配置与初始化 =================
st.set_page_config(
    page_title="员工职业健康体检报告在线查阅与签收平台",
    layout="centered",
    initial_sidebar_state="expanded",
)

# 注入华文宋体全局样式与隐藏默认元素
hide_streamlit_style = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    html, body, [class*="css"] {
        font-family: "华文宋体", SimSun, serif;
    }
    </style>
    """
st.markdown(hide_streamlit_style, unsafe_allow_html=True)


# 智能匹配员工专属 PDF 体检报告函数
def find_employee_pdf_report(folder, id_card, name):
  if not os.path.exists(folder):
    return None, None
  for filename in os.listdir(folder):
    if (
        id_card in filename
        and name in filename
        and filename.lower().endswith(".pdf")
    ):
      return os.path.join(folder, filename), filename
  for filename in os.listdir(folder):
    if id_card in filename and filename.lower().endswith(".pdf"):
      return os.path.join(folder, filename), filename
  return None, None


# ================= 2. 百度网盘自动上传函数（带 OAuth2 自动刷新） =================
def refresh_baidu_access_token():
  try:
    client_id = st.secrets.get("BAIDU_CLIENT_ID", "")
    client_secret = st.secrets.get("BAIDU_CLIENT_SECRET", "")
    refresh_token = st.secrets.get("BAIDU_REFRESH_TOKEN", "")
    if not client_id or not client_secret or not refresh_token:
      return None
    token_url = "https://pan.baidu.com/oauth/2.0/token"
    params = {
        "grant_type": "refresh_token",
        "refresh_token": refresh_token,
        "client_id": client_id,
        "client_secret": client_secret,
    }
    response = requests.get(token_url, params=params)
    res_data = response.json()
    return res_data.get("access_token")
  except Exception:
    return None


def upload_to_baidu_netdisk_with_auto_refresh(file_bytes, remote_filename):
  access_token = st.secrets.get("BAIDU_ACCESS_TOKEN", "")
  if not access_token:
    return (
        False,
        "未配置网盘凭证，文件已在本地生成并可通过网页下载/ZIP打包保存。",
    )

  sub_folder = "体检报告签收档案"
  target_path = f"/apps/慧瑞EHS合规档案/{sub_folder}/{remote_filename}"

  def send_upload_request(token):
    upload_url = f"https://pan.baidu.com/rest/2.0/xpan/file?method=upload&access_token={token}&path={target_path}&uploadid=&file=1"
    files = {"file": (remote_filename, file_bytes)}
    return requests.post(upload_url, files=files).json()

  result = send_upload_request(access_token)
  if "errno" in result and result["errno"] in [110, 111]:
    new_token = refresh_baidu_access_token()
    if new_token:
      result = send_upload_request(new_token)
    else:
      return False, "Token 已过期且自动刷新失败。"

  if "errno" in result and result["errno"] == 0:
    return True, f"成功同步至网盘：/apps/慧瑞EHS合规档案/{sub_folder}/"
  else:
    return False, f"网盘上传失败: {result.get('error_msg', '未知错误')}"


# ================= 3. 侧边栏：Logo、微信分享与说明 =================
with st.sidebar:
  try:
    st.image("logo.png", width=160)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=160,
    )

  st.markdown("### 📱 微信扫码与分享")
  st.write("已自动关联您的云端网址，二维码将实时更新供手机扫码填报。")

  app_url = st.text_input(
      "应用公网链接 (URL)", value="https://iway--technician.streamlit.app"
  )

  if app_url:
    qr = qrcode.make(app_url)
    img_buffer = io.BytesIO()
    qr.save(img_buffer, format="PNG")
    st.image(
        Image.open(img_buffer), caption="微信扫码快速查阅与签收", width=160
    )
    st.info(
        "💡 **提示**：将上方链接复制并发送至微信工作群，员工即可手机端完成体检报告签收。"
    )

  st.markdown("---")
  st.markdown("### 📂 管理员提示")
  st.write(
      "请将员工体检报告 **PDF 原件**放入 GitHub 仓库的 **体检报告/** 文件夹中。"
  )

# ================= 4. 主界面逻辑（Logo在左侧，主标题单独一行） =================
col_logo, col_title = st.columns([1, 6])
with col_logo:
  try:
    st.image("logo.png", width=110)
  except Exception:
    st.image(
        "https://upload.wikimedia.org/wikipedia/commons/thumb/c/c5/Ikea_logo.svg/800px-Ikea_logo.svg.png",
        width=110,
    )
with col_title:
  st.markdown("## 员工职业健康体检报告在线查阅与签收平台")

st.markdown(
    "请先准确输入您的姓名与18位身份证号，系统将自动检索并匹配您的专属 PDF"
    " 体检报告。查阅完毕后，请在底部完成手写签名与手写日期。"
)

# 基础信息录入
st.subheader("1. 员工身份核验")
col1, col2 = st.columns(2)
with col1:
  emp_name = st.text_input("员工姓名 (必填)：")
with col2:
  emp_id = st.text_input(
      "身份证号 (必填，须满18位)：",
      help="请输入标准的 18 位中国居民身份证号码",
  )

st.write("---")
st.markdown("### 📄 专属体检报告 PDF 查阅")

report_folder = "体检报告"
matched_path, matched_filename = None, None

id_pattern_check = re.compile(r"^\d{17}[\dXx]$")
if emp_name.strip() and id_pattern_check.match(emp_id.strip()):
  matched_path, matched_filename = find_employee_pdf_report(
      report_folder, emp_id.strip(), emp_name.strip()
  )

if matched_path and os.path.exists(matched_path):
  st.success(f"✅ 成功找到您的专属体检报告文件：【{matched_filename}】")
  with open(matched_path, "rb") as fr:
    report_bytes = fr.read()

  st.download_button(
      label=f"📥 点击下载并查阅《{matched_filename}》",
      data=report_bytes,
      file_name=matched_filename,
      mime="application/pdf",
      use_container_width=True,
  )
else:
  if emp_name.strip() or emp_id.strip():
    st.warning(
        "⚠️ 未在后台 '体检报告' 文件夹中检索到与您姓名及身份证匹配的 PDF"
        " 报告文件，请确认已由管理员上传。"
    )
  else:
    st.info("💡 请先在上方输入您的姓名和18位身份证号以加载您的体检报告。")

c_report = st.checkbox(
    "【须确认】本人已收到并查阅上述职业健康体检报告，已知悉体检结论及职业健康防护建议。"
)

# ================= 5. 手写签名与手写日期栏（并排双画布） =================
current_date_str = datetime.date.today().strftime("%Y年%m月%d日")

st.write("---")
st.subheader("✍️ 2. 员工手写签名与手写日期栏")
st.markdown(
    f"**请在左侧手写签名，并在右侧手写日期（注：当前系统日期为"
    f" {current_date_str}，请按此手写日期）：**"
)

col_sig, col_date = st.columns(2)
with col_sig:
  st.markdown("**手写签名：**")
  canvas_result = st_canvas(
      stroke_width=4,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_sig",
      return_image_data=True,
  )
with col_date:
  st.markdown("**手写日期栏（请手写当前日期）：**")
  canvas_date_result = st_canvas(
      stroke_width=3,
      stroke_color="#000000",
      background_color="#F8F9FA",
      height=200,
      width=320,
      drawing_mode="freedraw",
      key="canvas_date",
      return_image_data=True,
  )


# ================= 6. 辅助函数：生成 Word 格式的体检签收确认凭证 =================
def generate_medical_receipt_docx(
    employee_name, employee_id, report_name, sig_image_io, date_image_io
):
  doc = Document()

  p_title = doc.add_paragraph()
  run_t = p_title.add_run("【职业健康体检报告签收确认凭证】")
  run_t.font.name = "华文宋体"
  run_t.font.size = Pt(16)
  run_t.bold = True
  run_t.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  p_info = doc.add_paragraph()
  run_i = p_info.add_run(
      f"员工姓名: {employee_name}    身份证号: {employee_id}    签收时间:"
      f" {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
      f"关联体检报告文件: {report_name}\n"
      "本人已收到并查阅本人的职业健康体检报告，已知悉体检结论及各项健康指标与职业禁忌要求。"
  )
  run_i.font.name = "华文宋体"
  run_i.font.size = Pt(11)
  run_i.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")

  p_line = doc.add_paragraph(
      "--------------------------------------------------"
  )
  p_line.paragraph_format.space_before = Pt(5)
  p_line.paragraph_format.space_after = Pt(10)

  table = doc.add_table(rows=1, cols=2)
  table.autofit = False

  cell_sig = table.cell(0, 0)
  p1 = cell_sig.paragraphs[0]
  r1 = p1.add_run("员工手写亲笔签名：\n")
  r1.font.name = "华文宋体"
  r1.font.size = Pt(10.5)
  r1.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p1.add_run().add_picture(sig_image_io, width=Inches(1.8))
  sig_image_io.seek(0)

  cell_date = table.cell(0, 1)
  p2 = cell_date.paragraphs[0]
  r2 = p2.add_run("手写签署日期：\n")
  r2.font.name = "华文宋体"
  r2.font.size = Pt(10.5)
  r2.font.element.rPr.rFonts.set(qn("w:eastAsia"), "华文宋体")
  p2.add_run().add_picture(date_image_io, width=Inches(1.8))
  date_image_io.seek(0)

  buffer = io.BytesIO()
  doc.save(buffer)
  buffer.seek(0)
  return buffer


# ================= 7. 提交校验与生成档案 =================
if st.button(
    "📁 确认无误，一键签收体检报告并生成合规档案", use_container_width=True
):
  is_canvas_empty = canvas_result.image_data is None or (
      canvas_result.json_data is not None
      and len(canvas_result.json_data.get("objects", [])) == 0
  )
  is_date_empty = canvas_date_result.image_data is None or (
      canvas_date_result.json_data is not None
      and len(canvas_date_result.json_data.get("objects", [])) == 0
  )

  id_pattern = re.compile(r"^\d{17}[\dXx]$")

  if not emp_name.strip() or not emp_id.strip():
    st.error("❌ 拦截 : 请完整填写【员工姓名】与【身份证号】！")
  elif not id_pattern.match(emp_id.strip()):
    st.error(
        "❌ 拦截 : 身份证号必须为严格的 **18 位**数字（末尾可为大写 X）！"
    )
  elif not matched_path:
    st.error("❌ 拦截 : 未匹配到您的专属体检报告，无法完成签收！")
  elif not c_report:
    st.error("❌ 拦截 : 您必须勾选确认已查阅体检报告！")
  elif is_canvas_empty:
    st.warning("⚠️ 拦截 : 请在左侧画板完成手写签名后再提交。")
  elif is_date_empty:
    st.warning("⚠️ 拦截 : 请在右侧手写日期栏内完成手写日期后再提交！")
  else:
    st.success(
        "✅ 体检报告签收成功！系统已成功生成您的专属带签名 Word 合规确认凭证。"
    )

    signature_img = Image.fromarray(
        canvas_result.image_data.astype("uint8"), "RGBA"
    )
    sig_io = io.BytesIO()
    signature_img.save(sig_io, format="PNG")
    sig_io.seek(0)

    date_img = Image.fromarray(
        canvas_date_result.image_data.astype("uint8"), "RGBA"
    )
    date_io = io.BytesIO()
    date_img.save(date_io, format="PNG")
    date_io.seek(0)

    # 生成 Word 格式的体检签收凭证
    receipt_docx_buffer = generate_medical_receipt_docx(
        emp_name, emp_id, matched_filename, sig_io, date_io
    )

    # 动态打包 ZIP
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
      # 1. 放入员工原体检报告 PDF
      with open(matched_path, "rb") as fr:
        report_bytes = fr.read()
      zip_file.writestr(f"体检报告原件_{emp_name}_{matched_filename}", report_bytes)

      # 2. 放入带签名的签收确认凭证 Word (.docx)
      receipt_filename = f"体检报告签收确认凭证_{emp_name}_{emp_id[-4:]}.docx"
      zip_file.writestr(receipt_filename, receipt_docx_buffer.getvalue())

      # 3. 自动同步到百度网盘
      upload_to_baidu_netdisk_with_auto_refresh(
          receipt_docx_buffer.getvalue(), receipt_filename
      )

      # 4. 保存签名及日期原图
      img_byte_arr = io.BytesIO()
      signature_img.save(img_byte_arr, format="PNG")
      zip_file.writestr(
          f"手写签名原图_{emp_name}.png", img_byte_arr.getvalue()
      )

      date_byte_arr = io.BytesIO()
      date_img.save(date_byte_arr, format="PNG")
      zip_file.writestr(f"手写日期原图_{emp_name}.png", date_byte_arr.getvalue())

    zip_buffer.seek(0)

    st.markdown("---")
    st.success(
        "🎉 您的体检报告签收档案已打包完毕，点击下方按钮即可下载保存！"
    )

    col_d1, col_d2 = st.columns(2)
    with col_d1:
      st.download_button(
          label="📄 下载体检报告签收凭证 (.docx)",
          data=receipt_docx_buffer.getvalue(),
          file_name=f"体检报告签收确认凭证_{emp_name}.docx",
          mime=(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
          ),
          use_container_width=True,
      )
    with col_d2:
      st.download_button(
          label="📥 一键打包下载全部档案 (.ZIP)",
          data=zip_buffer,
          file_name=f"体检报告签收档案_{emp_name}.zip",
          mime="application/zip",
          use_container_width=True,
      )

    st.balloons()

# ================= 8. 底部版权与开发者声明 =================
st.markdown("---")
st.markdown(
    "<div style='text-align: center; color: gray; font-size: 14px;'>"
    "内部使用，严禁商业用途 | 开发者：陈野菲 Yefei"
    "</div>",
    unsafe_allow_html=True,
)
